"""DOCX paragraphs, tables, malformed packages and expansion guards."""

from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document

from src.data.ingestion import _parse_records, ingest_directory
from src.data.models import Status


def test_docx_order_and_tables(corpus, options):
    path = corpus / "valid.docx"
    doc = Document()
    doc.add_paragraph("Before")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left"
    table.cell(0, 1).text = "Right"
    doc.add_paragraph("After")
    doc.save(path)
    result = list(_parse_records(path, corpus, options))[0]
    assert result.text == "Before\nLeft\nRight\nAfter"
    assert result.metadata == {"paragraph_count": 4, "table_count": 1}


def test_empty_and_corrupt_docx(corpus, options):
    Document().save(corpus / "empty.docx")
    (corpus / "bad.docx").write_text("not a zip")
    result = ingest_directory(corpus, options)
    assert result.summary["no_extractable_text"] == 1
    assert result.summary["failed"] == 1


def test_expansion_guard(corpus, options):
    options.settings.values["ingestion"]["max_docx_uncompressed_mb"] = 1
    path = corpus / "expansion.docx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "a" * (1024**2 + 1))
    assert list(_parse_records(path, corpus, options))[0].status == Status.EXTRACTION_TOO_LARGE
